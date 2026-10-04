from docx import Document
from docx.shared import Inches

path = 'CSE307_Disk_Scheduling_Report.docx'
doc = Document(path)
for p in doc.paragraphs:
    if p.text.startswith('A disk scheduler decides'):
        p.add_run(' The definitions and trade-offs of SSTF, SCAN, and C-SCAN follow Arpaci-Dusseau and Arpaci-Dusseau [1].')
    if p.text.startswith('These are simulated requests'):
        p.add_run(' Linux also exposes contemporary I/O schedulers such as mq-deadline, BFQ, and Kyber, so this simulation should not be read as a Linux scheduler benchmark [2].')
    if p.text.startswith('Link: '):
        p.text = 'Repository URL: add your GitHub link after upload.'
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.startswith('____________________'):
                cell.text = '4 October 2026'

doc.add_heading('References', level=1)
doc.add_paragraph('[1] R. H. Arpaci-Dusseau and A. C. Arpaci-Dusseau, “Hard Disk Drives,” Operating Systems: Three Easy Pieces, version 1.10, 2023. https://pages.cs.wisc.edu/~remzi/OSTEP/file-disks.pdf')
doc.add_paragraph('[2] The Linux Kernel documentation, “Switching Scheduler.” https://docs.kernel.org/block/switching-sched.html (accessed 4 October 2026).')
doc.add_heading('Figures', level=1)
p = doc.add_paragraph('Figure 1. Seek-distance comparison on the same 300 held-out request batches.')
doc.add_picture('results/phase_comparison.png', width=Inches(5.7))
doc.add_paragraph('Figure 2. Workload shift: the same algorithms process each request batch.')
doc.add_picture('results/shift_timeline.png', width=Inches(5.7))
doc.save('CSE307_Disk_Scheduling_Final.docx')
